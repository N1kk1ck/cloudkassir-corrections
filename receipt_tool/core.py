"""Offline input validation and payload preparation. No network calls."""
from __future__ import annotations

import csv
import json
import os
import re
import uuid
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

import yaml

CENT = Decimal("0.01")
VAT_CODES = {None, 0, 5, 7, 10, 20, 22, 105, 107, 110, 120, 122}
RECEIPT_REQUIRED = {"entity_id", "receipt_id", "order_id", "receipt_type", "amount", "currency", "status"}
ITEM_REQUIRED = {"entity_id", "order_id", "line_id", "name", "quantity", "price", "amount"}


class InputError(ValueError):
    """Safe, actionable error that never includes configuration values."""


class UniqueKeyLoader(yaml.SafeLoader):
    pass


def _mapping(loader, node, deep=False):
    result = {}
    for key_node, value_node in node.value:
        key = loader.construct_object(key_node, deep=deep)
        if not isinstance(key, str) or key in result:
            raise InputError("YAML: ключи должны быть строками и не повторяться.")
        result[key] = loader.construct_object(value_node, deep=deep)
    return result


UniqueKeyLoader.add_constructor(yaml.resolver.BaseResolver.DEFAULT_MAPPING_TAG, _mapping)


def read_yaml(path: Path) -> dict:
    try:
        data = yaml.load(path.read_text(encoding="utf-8-sig"), Loader=UniqueKeyLoader)
    except (OSError, yaml.YAMLError) as exc:
        # YAML parser errors can contain the source line, including a secret.
        raise InputError("Не удалось прочитать YAML. Проверьте файл и синтаксис.") from None
    if not isinstance(data, dict):
        raise InputError("Верхний уровень YAML должен быть объектом.")
    return data


def integer(value: Any, allowed, field: str) -> int:
    if isinstance(value, bool) or str(value).strip() not in {str(n) for n in allowed}:
        raise InputError(f"Недопустимое значение поля {field}.")
    return int(value)


def number(value: Any, field: str, *, positive=True, cents=True) -> Decimal:
    try:
        result = Decimal(str(value).replace("\xa0", "").replace(" ", "").replace(",", "."))
    except (InvalidOperation, ValueError):
        raise InputError(f"Поле {field}: требуется число.") from None
    if not result.is_finite() or (result <= 0 if positive else result < 0):
        raise InputError(f"Поле {field}: требуется {'положительное' if positive else 'неотрицательное'} число.")
    if result > Decimal("1000000000000"):
        raise InputError(f"Поле {field}: число превышает поддерживаемый предел.")
    if cents and result != result.quantize(CENT):
        raise InputError(f"Поле {field}: допустимы не более двух знаков после запятой.")
    return result


def load_organizations(path: Path) -> dict[str, dict]:
    data = read_yaml(path).get("organizations")
    if not isinstance(data, dict) or not data:
        raise InputError("В YAML нужен непустой объект organizations.")
    result = {}
    for entity_id, organization in data.items():
        if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_-]{0,63}", entity_id):
            raise InputError("entity_id: используйте латинские буквы, цифры, дефис и подчёркивание.")
        if not isinstance(organization, dict):
            raise InputError(f"{entity_id}: настройки должны быть объектом.")
        fiscal = organization.get("fiscal")
        if not isinstance(fiscal, dict):
            raise InputError(f"{entity_id}: отсутствует fiscal.")
        for field in ["organization_inn", "payment_place", "payment_address", "correction_number", "correction_date"]:
            if not isinstance(fiscal.get(field), str) or not fiscal[field].strip():
                raise InputError(f"{entity_id}: заполните строковое поле fiscal.{field}.")
        if not re.fullmatch(r"\d{10}|\d{12}", fiscal["organization_inn"]):
            raise InputError(f"{entity_id}: ИНН должен содержать 10 или 12 цифр.")
        if fiscal.get("ffd_version") != "1.2":
            raise InputError(f"{entity_id}: поддерживается ФФД 1.2.")
        try:
            date.fromisoformat(fiscal["correction_date"])
        except ValueError:
            raise InputError(f"{entity_id}: correction_date должна иметь формат YYYY-MM-DD.") from None
        integer(fiscal.get("taxation_system"), range(6), "taxation_system")
        integer(fiscal.get("russia_time_zone"), range(1, 12), "russia_time_zone")
        integer(fiscal.get("item_payment_object", 1), range(1, 14), "item_payment_object")
        if "vat" not in fiscal or fiscal["vat"] not in VAT_CODES or isinstance(fiscal["vat"], bool):
            raise InputError(f"{entity_id}: явно задайте поддерживаемый код vat или null.")
        for field in ["demo", "fiscal_confirmed"]:
            if field in organization and not isinstance(organization[field], bool):
                raise InputError(f"{entity_id}: {field} должен быть true или false.")
        if not isinstance(organization.get("credentials", {}), dict):
            raise InputError(f"{entity_id}: credentials должен быть объектом.")
        result[entity_id] = organization
    return result


def credentials_for(organization: dict) -> tuple[str, str]:
    credentials = organization.get("credentials", {})
    values = []
    for key in ["public_id", "api_secret"]:
        env_name = credentials.get(key + "_env")
        if env_name is not None and not isinstance(env_name, str):
            raise InputError("Имя переменной окружения должно быть строкой.")
        value = os.environ.get(env_name, "") if env_name else credentials.get(key, "")
        if not isinstance(value, str) or not value.strip() or value.startswith(("YOUR_", "EXAMPLE_", "DEMO_")):
            raise InputError("Не заданы рабочие API-доступы организации; проверьте конфигурацию или окружение.")
        values.append(value.strip())
    return tuple(values)


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_table(path: Path, required: set[str], spec: dict | None = None) -> list[dict[str, str]]:
    spec = spec or {}
    columns = spec.get("columns", {})
    if not isinstance(columns, dict) or any(not isinstance(k, str) or not isinstance(v, str) for k, v in columns.items()):
        raise InputError("columns должен сопоставлять имена стандартных полей именам колонок.")
    if len(set(columns.values())) != len(columns):
        raise InputError("Одна исходная колонка не может обозначать несколько стандартных полей.")
    translations = spec.get("values", {})
    if not isinstance(translations, dict) or any(
        not isinstance(field, str) or not isinstance(values, dict)
        or any(not isinstance(k, str) or not isinstance(v, str) for k, v in values.items())
        for field, values in translations.items()
    ):
        raise InputError("values должен содержать строковые соответствия значений колонок.")
    header_row = integer(spec.get("header_row", 1), range(1, 1001), "header_row")
    workbook = None
    try:
        if path.suffix.lower() == ".csv":
            handle = path.open(encoding="utf-8-sig", newline="")
            try:
                sample = handle.read(8192)
                handle.seek(0)
                try:
                    dialect = csv.Sniffer().sniff("\n".join(sample.splitlines()[header_row - 1:]), delimiters=",;\t")
                except csv.Error:
                    dialect = csv.excel
                raw = list(csv.reader(handle, dialect))
            finally:
                handle.close()
        elif path.suffix.lower() == ".xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=True)
            if spec.get("sheet"):
                if spec["sheet"] not in workbook.sheetnames:
                    raise InputError("В XLSX не найден указанный лист.")
                sheet = workbook[spec["sheet"]]
            else:
                visible = [s for s in workbook.worksheets if s.sheet_state == "visible"]
                if len(visible) != 1:
                    raise InputError("В XLSX несколько видимых листов: задайте sheet в mapping.yaml.")
                sheet = visible[0]
            raw = list(sheet.iter_rows(values_only=True))
        else:
            raise InputError("Поддерживаются только CSV и XLSX.")
    except InputError:
        raise
    except Exception:
        raise InputError("Не удалось прочитать таблицу. Проверьте путь, формат и доступ к файлу.") from None
    finally:
        if workbook is not None:
            workbook.close()
    if len(raw) < header_row:
        raise InputError("Входная таблица пуста.")
    headers = [_cell(v).lstrip("\ufeff") for v in raw[header_row - 1]]
    nonempty = [h for h in headers if h]
    if len(nonempty) != len(set(nonempty)):
        raise InputError("В таблице повторяются заголовки.")
    reverse = {source: standard for standard, source in columns.items()}
    mapped = [reverse.get(h, h) for h in headers]
    if len([h for h in mapped if h]) != len({h for h in mapped if h}):
        raise InputError("После сопоставления повторяются имена полей.")
    missing = required - set(mapped)
    if missing:
        raise InputError("Не найдены обязательные колонки: " + ", ".join(sorted(missing)))
    result = []
    for row_number, values in enumerate(raw[header_row:], header_row + 1):
        if not any(_cell(v) for v in values):
            continue
        if any(_cell(v) for v in values[len(headers):]):
            raise InputError(f"Строка {row_number}: больше значений, чем заголовков.")
        row = {name: _cell(values[i]) if i < len(values) else "" for i, name in enumerate(mapped) if name}
        for field, lookup in translations.items():
            if field in row:
                row[field] = lookup.get(row[field], row[field])
        row["_row"] = str(row_number)
        result.append(row)
    if not result:
        raise InputError("В таблице нет строк данных.")
    return result


def contact_for(receipt: dict) -> str:
    email = receipt.get("email", "").strip()
    if email:
        if re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
            return email
        raise InputError("Некорректный email.")
    phone = re.sub(r"[\s()+-]", "", receipt.get("phone", ""))
    if len(phone) == 10 and phone.isdigit():
        phone = "7" + phone
    if len(phone) == 11 and phone.startswith("8"):
        phone = "7" + phone[1:]
    if not re.fullmatch(r"[1-9]\d{9,14}", phone):
        raise InputError("Нужен email или корректный телефон покупателя.")
    return "+" + phone


def api_number(value: Decimal) -> int | float:
    return int(value) if value == value.to_integral_value() else float(value)


def receipt_key(receipt: dict) -> str:
    return f"{receipt['entity_id']}:{receipt['receipt_id']}:{receipt['receipt_type']}"


def request_id(receipt: dict) -> str:
    # JSON encoding prevents separator collisions in user-supplied identifiers.
    source = json.dumps(["cloudkassir-corrections-v1", receipt["entity_id"], receipt["receipt_id"], receipt["receipt_type"]], ensure_ascii=False)
    return str(uuid.uuid5(uuid.NAMESPACE_URL, source))


def build_payload(receipt: dict, lines: list[dict], organization: dict) -> dict:
    if receipt["receipt_type"] not in {"income", "income_return"}:
        raise InputError("receipt_type: допустимы income и income_return.")
    if receipt["currency"] != "RUB":
        raise InputError("Поддерживаются только платежи в RUB.")
    if receipt["status"].lower() not in {"completed", "paid"}:
        raise InputError("Операция должна иметь статус completed или paid.")
    if receipt["receipt_type"] == "income_return" and receipt.get("items_confirmed", "").lower() != "true":
        raise InputError("Для возврата подтвердите состав: items_confirmed=true.")
    total = number(receipt["amount"], "amount")
    if not lines:
        raise InputError("Не найден состав чека.")
    fiscal = organization["fiscal"]
    items = []
    for line in lines:
        name = line["name"]
        if not name or len(name) > 128:
            raise InputError("Название позиции должно содержать от 1 до 128 символов.")
        quantity = number(line["quantity"], "quantity")
        price = number(line["price"], "price", positive=False)
        amount = number(line["amount"], "item.amount", positive=False)
        if amount > (quantity * price).quantize(CENT, rounding=ROUND_HALF_UP):
            raise InputError("Сумма позиции превышает количество × цену.")
        vat = fiscal["vat"]
        if line.get("vat", ""):
            vat = None if line["vat"].lower() in {"null", "none"} else integer(line["vat"], VAT_CODES - {None}, "vat")
        items.append({
            "label": name, "quantity": api_number(quantity), "price": api_number(price),
            "amount": api_number(amount), "vat": vat,
            "object": integer(line.get("object") or fiscal.get("item_payment_object", 1), range(1, 14), "object"),
            "method": 4,
        })
    if sum(number(line["amount"], "item.amount", positive=False) for line in lines) != total:
        raise InputError("Сумма позиций не совпадает с суммой чека до копейки.")
    data = {
        "organizationInn": fiscal["organization_inn"],
        "paymentPlace": fiscal["payment_place"],
        "paymentAddress": fiscal["payment_address"],
        "taxationSystem": int(fiscal["taxation_system"]),
        "correctionType": 0,
        "correctionReceiptType": 1 if receipt["receipt_type"] == "income" else 2,
        "causeCorrection": {"correctionNumber": fiscal["correction_number"], "correctionDate": fiscal["correction_date"]},
        "items": items,
        "amounts": {"electronic": api_number(total), "cash": 0, "advancePayment": 0, "credit": 0, "provision": 0},
        "isInternetPayment": True,
        "russiaTimeZone": int(fiscal["russia_time_zone"]),
        "customerContactAddress": contact_for(receipt),
    }
    if fiscal.get("cashier_name"):
        data["cashierName"] = fiscal["cashier_name"]
    return {"correctionReceiptData": data}


@dataclass
class Preparation:
    ready: list[dict]
    review: list[dict]
    unused_items: int


def prepare(organizations: dict, receipts: list[dict], items: list[dict]) -> Preparation:
    seen = set()
    order_counts = Counter()
    for receipt in receipts:
        if receipt["entity_id"] not in organizations:
            raise InputError("В списке чеков есть незарегистрированный entity_id.")
        if not receipt["receipt_id"] or not receipt["order_id"]:
            raise InputError("receipt_id и order_id не могут быть пустыми.")
        key = (receipt["entity_id"], receipt["receipt_id"])
        if key in seen:
            raise InputError("В списке чеков повторяется пара entity_id + receipt_id.")
        seen.add(key)
        order_counts[(receipt["entity_id"], receipt["order_id"])] += 1
    by_order = defaultdict(list)
    seen_lines = set()
    for line in items:
        if line["entity_id"] not in organizations or not line["order_id"] or not line["line_id"]:
            raise InputError("У позиции должны быть зарегистрированный entity_id, order_id и line_id.")
        key = (line["entity_id"], line["order_id"], line.get("receipt_id", ""), line["line_id"])
        if key in seen_lines:
            raise InputError("В составе заказа повторяется line_id.")
        seen_lines.add(key)
        by_order[key[:2]].append(line)
    ready, review, used = [], [], set()
    for receipt in receipts:
        candidates = by_order[(receipt["entity_id"], receipt["order_id"])]
        scoped = [line for line in candidates if line.get("receipt_id", "") == receipt["receipt_id"]]
        unscoped = [line for line in candidates if not line.get("receipt_id", "")]
        lines = scoped if scoped else unscoped
        try:
            if scoped and unscoped:
                raise InputError("Смешан состав по заказу и по отдельному чеку.")
            if unscoped and order_counts[(receipt["entity_id"], receipt["order_id"])] > 1:
                raise InputError("Несколько чеков на заказ: у каждой позиции нужен receipt_id.")
            payload = build_payload(receipt, lines, organizations[receipt["entity_id"]])
        except InputError as exc:
            review.append({"entity_id": receipt["entity_id"], "receipt_id": receipt["receipt_id"], "order_id": receipt["order_id"], "reason": str(exc)})
            continue
        used.update(id(line) for line in lines)
        ready.append({"entity_id": receipt["entity_id"], "receipt_id": receipt["receipt_id"], "order_id": receipt["order_id"], "receipt_type": receipt["receipt_type"], "receipt_key": receipt_key(receipt), "request_id": request_id(receipt), "payload": payload})
    return Preparation(ready, review, sum(id(line) not in used for line in items))


def load_inputs(config: Path, receipts: Path, items: Path, mapping: Path | None = None):
    spec = read_yaml(mapping) if mapping else {}
    for table in ["receipts", "items"]:
        if not isinstance(spec.get(table, {}), dict):
            raise InputError("Настройка таблицы в mapping.yaml должна быть объектом.")
    organizations = load_organizations(config)
    return organizations, prepare(organizations, read_table(receipts, RECEIPT_REQUIRED, spec.get("receipts")), read_table(items, ITEM_REQUIRED, spec.get("items")))


def csv_safe(value: Any) -> Any:
    text = str(value)
    if text.startswith(("=", "+", "-", "@", "\t", "\r")):
        try:
            if not Decimal(text).is_finite():
                return "'" + text
        except InvalidOperation:
            return "'" + text
    return value


def write_csv(path: Path, fields: list[str], rows: list[dict]):
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows({key: csv_safe(row.get(key, "")) for key in fields} for row in rows)


def write_preparation(output: Path, result: Preparation) -> dict:
    output.mkdir(parents=True, exist_ok=True)
    preview = output / "payload_preview.jsonl"
    temp = preview.with_suffix(".tmp")
    with temp.open("w", encoding="utf-8") as handle:
        for row in result.ready:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")
    temp.replace(preview)
    write_csv(output / "ready.csv", ["entity_id", "receipt_id", "order_id", "receipt_type", "request_id"], result.ready)
    write_csv(output / "review.csv", ["entity_id", "receipt_id", "order_id", "reason"], result.review)
    summary = {"receipts_total": len(result.ready) + len(result.review), "ready": len(result.ready), "review": len(result.review), "unused_item_rows": result.unused_items}
    (output / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return summary
